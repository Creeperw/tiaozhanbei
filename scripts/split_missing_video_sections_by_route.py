from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = ["路径内顺序", "书", "章节名", "小节名", "视频链接"]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="按学习路径顺序拆分待补视频小节。")
    parser.add_argument("--app-root", type=Path, required=True)
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--chapter-root", type=Path, required=True)
    parser.add_argument("--target-catalog", type=Path, required=True)
    parser.add_argument("--textbook-catalog", type=Path, required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--first-output", type=Path, required=True)
    parser.add_argument("--rest-output", type=Path, required=True)
    parser.add_argument("--limit", type=int, default=25)
    return parser.parse_args()


def style_sheet(sheet, table_name: str) -> None:
    header_fill = PatternFill("solid", fgColor="14865F")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="thin", color="D9E9E2"))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="E8F0EC"))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{max(sheet.max_row, 1)}"
    for column, width in zip("ABCDE", (12, 24, 34, 44, 48)):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 26
    sheet.sheet_view.showGridLines = False
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=f"A1:E{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)


def workbook_for_routes(route_rows: list[tuple[dict, list[tuple]]]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = workbook.create_sheet("汇总")
    summary.append(["学习路径", "教材范围", "教材数", "待补视频小节数"])
    for index, (route, rows) in enumerate(route_rows, start=1):
        title = route["name"][:31]
        sheet = workbook.create_sheet(title)
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
        style_sheet(sheet, f"QualificationRoute{index}Rows")
        summary.append(
            [route["official_name"], route["range_label"], route["book_count"], len(rows)]
        )
    summary.freeze_panes = "A2"
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 20
    summary.column_dimensions["C"].width = 14
    summary.column_dimensions["D"].width = 20
    summary.sheet_view.showGridLines = False
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor="175C45")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    return workbook


def main() -> None:
    args = parse_args()
    sys.path.insert(0, str(args.app_root.resolve()))
    from APP.backend.knowledge_atlas_service import KnowledgeAtlasStore

    store = KnowledgeAtlasStore(
        args.data_root,
        chapter_root=args.chapter_root,
        enabled=True,
    )
    store.ensure_hierarchy()

    source_book = load_workbook(args.source, read_only=True, data_only=True)
    source_sheet = source_book["待补视频"]
    missing_by_book: dict[str, list[tuple[str, str, str]]] = {}
    for book, chapter, section, _video_link in source_sheet.iter_rows(
        min_row=2, values_only=True
    ):
        missing_by_book.setdefault(str(book), []).append(
            (str(chapter), str(section), "")
        )

    all_books = sorted(
        store.chapters_by_book,
        key=lambda name: (store.book_order.get(name, 999999), name),
    )
    target_payload = json.loads(args.target_catalog.read_text(encoding="utf-8-sig"))
    textbook_payload = json.loads(args.textbook_catalog.read_text(encoding="utf-8-sig"))
    textbook_routes = {
        route["route_id"]: route for route in textbook_payload["routes"]
    }
    first_routes: list[tuple[dict, list[tuple]]] = []
    selected_books: set[str] = set()
    for target in target_payload["items"]:
        route = textbook_routes[target["textbook_route_id"]]
        ordered_books: list[str] = []
        for stage in sorted(route["stages"], key=lambda row: row["order"]):
            for decorated_book in stage["books"]:
                book = str(decorated_book).strip().removeprefix("《").removesuffix("》")
                if book not in ordered_books:
                    ordered_books.append(book)
        first_books = ordered_books[: args.limit]
        selected_books.update(book for book in first_books if book in store.chapters_by_book)

        def rows_for(books: list[str]) -> list[tuple]:
            rows: list[tuple] = []
            order_by_book = {book: index for index, book in enumerate(ordered_books, 1)}
            for book in books:
                for chapter, section, video_link in missing_by_book.get(book, []):
                    rows.append(
                        (order_by_book[book], book, chapter, section, video_link)
                    )
            return rows

        first_meta = {
            **target,
            "name": target["official_name"],
            "range_label": f"第1—{min(args.limit, len(ordered_books))}本",
            "book_count": len(first_books),
        }
        first_routes.append((first_meta, rows_for(first_books)))

    remaining_books = [book for book in all_books if book not in selected_books]
    remaining_rows: list[tuple] = []
    for order, book in enumerate(remaining_books, 1):
        for chapter, section, video_link in missing_by_book.get(book, []):
            remaining_rows.append((order, book, chapter, section, video_link))

    args.first_output.parent.mkdir(parents=True, exist_ok=True)
    args.rest_output.parent.mkdir(parents=True, exist_ok=True)
    first_workbook = workbook_for_routes(first_routes)
    rest_workbook = Workbook()
    rest_sheet = rest_workbook.active
    rest_sheet.title = "其余教材"
    rest_sheet.append(HEADERS)
    for row in remaining_rows:
        rest_sheet.append(row)
    style_sheet(rest_sheet, "RemainingTextbookRows")
    rest_summary = rest_workbook.create_sheet("统计说明", 0)
    rest_summary.append(["范围", "教材数", "待补视频小节数"])
    rest_summary.append(["不属于5个当前考试类别前25本的教材", len(remaining_books), len(remaining_rows)])
    rest_summary.column_dimensions["A"].width = 48
    rest_summary.column_dimensions["B"].width = 14
    rest_summary.column_dimensions["C"].width = 20
    for cell in rest_summary[1]:
        cell.fill = PatternFill("solid", fgColor="175C45")
        cell.font = Font(color="FFFFFF", bold=True)
    first_workbook.save(args.first_output)
    rest_workbook.save(args.rest_output)

    for path in (args.first_output, args.rest_output):
        checked = load_workbook(path, read_only=True, data_only=True)
        assert checked.sheetnames

    print(f"FIRST_OUTPUT={args.first_output.resolve()}")
    print(f"REST_OUTPUT={args.rest_output.resolve()}")
    for first_meta, first_rows in first_routes:
        print(
            f"{first_meta['official_name']}: books={first_meta['book_count']}, "
            f"rows={len(first_rows)}"
        )
    print(f"REMAINING: books={len(remaining_books)}, rows={len(remaining_rows)}")


if __name__ == "__main__":
    main()
