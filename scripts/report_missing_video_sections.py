from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="导出既无小节视频、也无任何知识点时间戳的小节。",
    )
    parser.add_argument("--app-root", type=Path, required=True)
    parser.add_argument("--data-root", type=Path, required=True)
    parser.add_argument("--video-root", type=Path, required=True)
    parser.add_argument("--chapter-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def knowledge_point_id(row: dict) -> str:
    return str(row.get("kp_id") or row.get("id") or "")


def main() -> None:
    args = parse_args()
    sys.path.insert(0, str(args.app_root.resolve()))
    from APP.backend.knowledge_atlas_service import KnowledgeAtlasStore

    store = KnowledgeAtlasStore(
        args.data_root,
        video_root=args.video_root,
        chapter_root=args.chapter_root,
        enabled=True,
    )
    store.ensure_hierarchy()
    store.ensure_videos()

    all_sections: list[dict] = []
    missing_rows: list[tuple[str, str, str, str]] = []
    for book in sorted(
        store.chapters_by_book,
        key=lambda name: (store.book_order.get(name, 999999), name),
    ):
        for chapter in store.chapters_by_book[book]:
            for section in chapter["sections"]:
                all_sections.append(section)
                has_section_video = bool(
                    store.section_videos_by_section.get(section["id"])
                )
                has_knowledge_timestamp = any(
                    bool(store.videos_by_kp.get(knowledge_point_id(kp)))
                    for kp in store.kps_by_section.get(section["id"], [])
                )
                if not has_section_video and not has_knowledge_timestamp:
                    missing_rows.append(
                        (book, chapter["name"], section["name"], "")
                    )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "待补视频"
    sheet.append(["书", "章节名", "小节名", "视频链接"])
    for row in missing_rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="14865F")
    header_font = Font(color="FFFFFF", bold=True)
    header_border = Border(bottom=Side(style="thin", color="D9E9E2"))
    row_border = Border(bottom=Side(style="hair", color="E8F0EC"))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = row_border

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{max(sheet.max_row, 1)}"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 34
    sheet.column_dimensions["C"].width = 44
    sheet.column_dimensions["D"].width = 48
    sheet.row_dimensions[1].height = 26
    sheet.sheet_view.showGridLines = False
    if missing_rows:
        table = Table(displayName="MissingVideoSections", ref=f"A1:D{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    no_section_but_kp_timestamp = sum(
        not bool(store.section_videos_by_section.get(section["id"]))
        and any(
            bool(store.videos_by_kp.get(knowledge_point_id(kp)))
            for kp in store.kps_by_section.get(section["id"], [])
        )
        for section in all_sections
    )
    summary = workbook.create_sheet("统计说明")
    summary_rows = [
        ("统计口径", "同时满足：该小节无小节完整视频，且该小节下所有知识点均无视频时间戳"),
        ("教材数", len(store.chapters_by_book)),
        ("全部小节数", len(all_sections)),
        ("待补视频小节数", len(missing_rows)),
        (
            "已有小节视频的小节数",
            sum(
                bool(store.section_videos_by_section.get(section["id"]))
                for section in all_sections
            ),
        ),
        (
            "无小节视频但至少有一个知识点时间戳的小节数",
            no_section_but_kp_timestamp,
        ),
        ("知识点时间戳数据源", str(args.video_root.resolve())),
        ("章节、小节及小节视频数据源", str(args.chapter_root.resolve())),
    ]
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 110
    summary.sheet_view.showGridLines = False
    for cell in summary["A"]:
        cell.font = Font(bold=True, color="175C45")
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.save(args.output)

    checked = load_workbook(args.output, read_only=True, data_only=True)
    assert checked["待补视频"].max_row == len(missing_rows) + 1
    if missing_rows:
        assert checked["待补视频"]["D2"].value is None

    print(f"OUTPUT={args.output.resolve()}")
    print(f"BOOKS={len(store.chapters_by_book)}")
    print(f"ALL_SECTIONS={len(all_sections)}")
    print(f"MISSING={len(missing_rows)}")
    print(
        "SECTION_VIDEO_RESOLVED="
        f"{store._section_video_metrics['resolved_rows']}"
    )
    print(f"KP_WITH_TIMESTAMP={sum(bool(rows) for rows in store.videos_by_kp.values())}")
    print(f"TOP_BOOKS={Counter(row[0] for row in missing_rows).most_common(10)}")


if __name__ == "__main__":
    main()
