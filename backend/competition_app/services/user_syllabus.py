from __future__ import annotations

import asyncio
import base64
from concurrent.futures import ThreadPoolExecutor
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import textwrap
import time
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

import httpx
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont, ImageSequence
from pypdf import PdfReader

USER_SYLLABUS_UNSUPPORTED_FILE = "USER_SYLLABUS_UNSUPPORTED_FILE"
USER_SYLLABUS_EXTRACTION_FAILED = "USER_SYLLABUS_EXTRACTION_FAILED"
USER_SYLLABUS_INVALID_STRUCTURE = "USER_SYLLABUS_INVALID_STRUCTURE"
USER_SYLLABUS_NOT_FOUND = "USER_SYLLABUS_NOT_FOUND"
SUPPORTED_SUFFIXES = {
    ".pdf", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
    ".doc", ".docx", ".xls", ".xlsx", ".md", ".markdown", ".txt", ".csv",
}


class UserSyllabusError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code, self.message = code, message


def _split_text(text: str, size: int) -> list[str]:
    if len(text) <= size:
        return [text]
    chunks: list[str] = []
    current = text
    while len(current) > size:
        cut = current.rfind("\n", 0, size)
        if cut <= 0:
            cut = size
        chunks.append(current[:cut].strip())
        current = current[cut:].strip()
    if current:
        chunks.append(current)
    return chunks


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_owner(value: str) -> str:
    owner = str(value or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9_.@-]{1,160}", owner):
        raise ValueError("owner_id 格式无效")
    return owner


def _json_object(raw: Any) -> dict[str, Any]:
    if isinstance(raw, dict):
        return raw
    if isinstance(raw, list):
        raw = "".join(
            str(item.get("text") or item.get("content") or "") if isinstance(item, dict) else str(item)
            for item in raw
        )
    text = str(raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I)
    try:
        value = json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start < 0 or end <= start:
            raise UserSyllabusError(USER_SYLLABUS_INVALID_STRUCTURE, "多模态模型未返回有效 JSON")
        try:
            value = json.loads(text[start : end + 1])
        except json.JSONDecodeError as exc:
            snippet = text[start : end + 1].replace("\n", " ")[:200]
            raise UserSyllabusError(
                USER_SYLLABUS_INVALID_STRUCTURE,
                f"多模态模型返回的 JSON 无法解析（片段：{snippet}）") from exc
    if not isinstance(value, dict):
        raise UserSyllabusError(USER_SYLLABUS_INVALID_STRUCTURE, "考纲结构必须是 JSON 对象")
    return value


class UserSyllabusService:
    max_upload_bytes = 256 * 1024 * 1024
    image_batch_size = 6

    def __init__(self, runtime_root: Path, *, chat_base_url: str, chat_model: str,
                 chat_api_key: str, timeout_seconds: float = 120.0,
                 mineru_token: str | None = None,
                 mineru_pipeline_root: Path | None = None,
                 knowledge_resolver: Any | None = None,
                 vector_matcher: Any | None = None) -> None:
        self.root = Path(runtime_root) / "user_syllabi"
        self.root.mkdir(parents=True, exist_ok=True)
        self.chat_base_url = str(chat_base_url).rstrip("/")
        self.chat_model = str(chat_model).strip()
        self.chat_api_key = str(chat_api_key).strip()
        self.timeout_seconds = float(timeout_seconds)
        self.mineru_token = str(mineru_token or "").strip()
        self.mineru_pipeline_root = (
            Path(mineru_pipeline_root).resolve() if mineru_pipeline_root else None
        )
        self.knowledge_resolver = knowledge_resolver
        self.vector_matcher = vector_matcher

    async def import_file(self, owner_id: str, filename: str, content: bytes, *,
                          title: str = "", subject: str = "", exam_type: str = "",
                          syllabus_id: str | None = None) -> dict[str, Any]:
        owner = _safe_owner(owner_id)
        suffix = Path(filename or "").suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise UserSyllabusError(USER_SYLLABUS_UNSUPPORTED_FILE, "不支持该考纲文件格式")
        if not content:
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "考纲文件不能为空")
        if len(content) > self.max_upload_bytes:
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "考纲文件不能超过 256 MB")
        if not self.chat_base_url or not self.chat_model or not self.chat_api_key:
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "多模态模型配置不完整")
        current_id = syllabus_id or f"USY_{uuid4().hex}"
        run_dir = self._syllabus_dir(owner, current_id)
        run_dir.mkdir(parents=True, exist_ok=True)
        source = run_dir / f"source{suffix}"
        source.write_bytes(content)
        previous = self._read_json(run_dir / "manifest.json", {})
        manifest = {
            "syllabus_id": current_id, "owner_user_id": owner, "scope": "personal",
            "origin": "user_upload", "title": str(title or previous.get("title") or Path(filename).stem).strip(),
            "subject": str(subject or previous.get("subject") or "").strip(),
            "exam_type": str(exam_type or previous.get("exam_type") or "").strip(),
            "source_filename": Path(filename).name, "source_suffix": suffix,
            "processing_status": "processing", "is_active": bool(previous.get("is_active")),
            "created_at": previous.get("created_at") or _now(), "updated_at": _now(), "error": None,
        }
        self._write_json(run_dir / "manifest.json", manifest)
        try:
            fragments: list[dict[str, Any]] = []
            pages: list[tuple[int, Path]] = []
            parser_label = "multimodal_model_only"
            if suffix == ".pdf" and self.mineru_token and self.mineru_pipeline_root:
                try:
                    markdown = await asyncio.to_thread(
                        self._mineru_markdown, source, run_dir / "mineru"
                    )
                    fragments = await self._extract_markdown_batches(markdown)
                    parser_label = "mineru"
                except UserSyllabusError:
                    fragments = []
            if not fragments:
                pages = self._render_input(source, run_dir / "rendered_pages")
                if not pages:
                    raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "考纲没有可识别内容")
                fragments = [
                    await self._extract_batch(pages[offset:offset + self.image_batch_size])
                    for offset in range(0, len(pages), self.image_batch_size)
                ]
            raw = await self._merge_fragments(fragments)
            structured, requirements = self._normalize_structure(
                raw, current_id, manifest["title"], manifest["subject"], manifest["exam_type"]
            )
            mappings = await self._map_requirements(requirements, structured)
            manifest.update({
                "title": structured["title"], "subject": structured["subject"],
                "exam_type": structured["exam_type"], "processing_status": "success",
                "page_count": len(pages) or max(
                    (page for row in requirements for page in row.get("source_pages") or []),
                    default=0,
                ),
                "section_count": len(structured["sections"]),
                "requirement_count": len(requirements),
                "matched_requirement_count": sum(row["match_status"] == "matched" for row in mappings),
                "updated_at": _now(), "error": None,
            })
            self._write_json(run_dir / "structured.json", structured)
            self._write_jsonl(run_dir / "requirements.jsonl", requirements)
            self._write_jsonl(run_dir / "mappings.jsonl", mappings)
            self._write_json(run_dir / "extraction_report.json", {
                "model": self.chat_model, "semantic_parser": parser_label,
                "rendered_page_count": len(pages), "fragment_count": len(fragments), "completed_at": _now(),
            })
            self._write_json(run_dir / "manifest.json", manifest)
            self.activate(owner, current_id)
            return self.get(owner, current_id)
        except UserSyllabusError as exc:
            manifest.update({"processing_status": "failed", "updated_at": _now(),
                             "error": {"code": exc.code, "message": exc.message}})
            self._write_json(run_dir / "manifest.json", manifest)
            raise
        except Exception as exc:
            error = UserSyllabusError(
                USER_SYLLABUS_EXTRACTION_FAILED, f"考纲处理失败：{type(exc).__name__}: {exc}"
            )
            manifest.update({"processing_status": "failed", "updated_at": _now(),
                             "error": {"code": error.code, "message": error.message}})
            self._write_json(run_dir / "manifest.json", manifest)
            raise error from exc

    async def reprocess(self, owner_id: str, syllabus_id: str) -> dict[str, Any]:
        owner = _safe_owner(owner_id)
        run_dir = self._require_dir(owner, syllabus_id)
        manifest = self._read_json(run_dir / "manifest.json", {})
        source = next((path for path in run_dir.glob("source.*") if path.is_file()), None)
        if source is None:
            raise UserSyllabusError(USER_SYLLABUS_NOT_FOUND, "考纲源文件不存在")
        return await self.import_file(owner, str(manifest.get("source_filename") or source.name),
            source.read_bytes(), title=str(manifest.get("title") or ""),
            subject=str(manifest.get("subject") or ""), exam_type=str(manifest.get("exam_type") or ""),
            syllabus_id=syllabus_id)

    def list(self, owner_id: str) -> list[dict[str, Any]]:
        owner_dir = self.root / _safe_owner(owner_id)
        if not owner_dir.is_dir():
            return []
        rows = [self._read_json(path / "manifest.json", {}) for path in owner_dir.iterdir()
                if path.is_dir() and (path / "manifest.json").is_file()]
        return sorted((row for row in rows if row.get("syllabus_id")),
                      key=lambda row: str(row.get("updated_at") or ""), reverse=True)

    def get(self, owner_id: str, syllabus_id: str) -> dict[str, Any]:
        run_dir = self._require_dir(_safe_owner(owner_id), syllabus_id)
        structured = self._read_json(run_dir / "structured.json", None)
        return {
            "manifest": self._read_json(run_dir / "manifest.json", {}),
            "structured": structured,
            "mappings": self._read_jsonl(run_dir / "mappings.jsonl"),
        }

    def requirements(self, owner_id: str, syllabus_id: str) -> list[dict[str, Any]]:
        return self._read_jsonl(self._require_dir(_safe_owner(owner_id), syllabus_id) / "requirements.jsonl")

    def mappings(self, owner_id: str, syllabus_id: str) -> list[dict[str, Any]]:
        return self._read_jsonl(self._require_dir(_safe_owner(owner_id), syllabus_id) / "mappings.jsonl")

    def activate(self, owner_id: str, syllabus_id: str) -> dict[str, Any]:
        owner = _safe_owner(owner_id)
        target = self._require_dir(owner, syllabus_id)
        if self._read_json(target / "manifest.json", {}).get("processing_status") != "success":
            raise UserSyllabusError(USER_SYLLABUS_INVALID_STRUCTURE, "只有处理成功的考纲可以激活")
        for path in (self.root / owner).iterdir():
            manifest_path = path / "manifest.json"
            if not manifest_path.is_file():
                continue
            manifest = self._read_json(manifest_path, {})
            active = path.name == syllabus_id
            if bool(manifest.get("is_active")) != active:
                manifest.update({"is_active": active, "updated_at": _now()})
                self._write_json(manifest_path, manifest)
        return self.get(owner, syllabus_id)

    def delete(self, owner_id: str, syllabus_id: str) -> None:
        owner = _safe_owner(owner_id)
        target = self._require_dir(owner, syllabus_id).resolve()
        if target.parent != (self.root / owner).resolve():
            raise UserSyllabusError(USER_SYLLABUS_NOT_FOUND, "考纲不存在")
        was_active = bool(self._read_json(target / "manifest.json", {}).get("is_active"))
        shutil.rmtree(target)
        if was_active:
            candidates = [row for row in self.list(owner) if row.get("processing_status") == "success"]
            if candidates:
                self.activate(owner, str(candidates[0]["syllabus_id"]))

    def load_context(self, owner_id: str, user_request: str) -> dict[str, Any]:
        active = next((row for row in self.list(owner_id)
                       if row.get("is_active") and row.get("processing_status") == "success"), None)
        if not active:
            return {}
        syllabus_id = str(active["syllabus_id"])
        requirements = self.requirements(owner_id, syllabus_id)
        mappings = {row.get("requirement_id"): row for row in self.mappings(owner_id, syllabus_id)}
        request_text = re.sub(r"\s+", "", str(user_request or "")).lower()
        identity_hit = any(re.sub(r"\s+", "", str(value or "")).lower() in request_text
                           for value in (active.get("title"), active.get("subject"), active.get("exam_type"))
                           if len(re.sub(r"\s+", "", str(value or ""))) >= 2)
        exam_intent = any(word in request_text for word in
                          ("考纲", "考试", "期末", "练题", "出题", "题目", "复习", "讲解", "知识点"))
        scored = []
        for requirement in requirements:
            text = "".join(str(requirement.get(key) or "") for key in ("title", "details", "section_title"))
            terms = set(re.findall(r"[\u4e00-\u9fff]{2,8}|[a-z0-9_]{2,}", text.lower()))
            score = sum(term in request_text for term in terms)
            if score:
                scored.append((score, requirement))
        if not identity_hit and not exam_intent and not scored:
            return {}
        selected = [row for _, row in sorted(scored, key=lambda pair: -pair[0])[:24]] or requirements[:24]
        selected_payload, kp_rows, seen = [], [], set()
        for requirement in selected:
            mapping = mappings.get(requirement.get("requirement_id"), {})
            selected_payload.append({**requirement, "mapping": mapping})
            kp_id = mapping.get("kp_id")
            if mapping.get("match_status") == "matched" and kp_id not in seen:
                seen.add(kp_id)
                kp_rows.append({"kp_id": kp_id, "kp_name": mapping.get("kp_name"),
                                "confidence": mapping.get("confidence"),
                                "requirement_id": requirement.get("requirement_id")})
        return {
            "user_syllabus": {"syllabus_id": syllabus_id, "title": active.get("title"),
                              "subject": active.get("subject"), "exam_type": active.get("exam_type"),
                              "scope": "personal", "origin": "user_upload"},
            "syllabus_requirements": selected_payload,
            "syllabus_knowledge_points": kp_rows,
        }

    async def _extract_batch(self, pages: list[tuple[int, Path]]) -> dict[str, Any]:
        content: list[dict[str, Any]] = [{"type": "text", "text": (
            "你是考试考纲结构化解析器。以下图片是同一份考纲的连续页面或虚拟页面。"
            "只根据图片内容识别，不得补写图片中不存在的要求。返回严格 JSON："
            "{\"document_title\":\"\",\"subject\":\"\",\"exam_type\":\"\","
            "\"sections\":[{\"title\":\"章节/模块名\",\"requirements\":[{"
            "\"title\":\"具体考核要求\",\"mastery_level\":\"掌握/熟悉/了解/未注明\","
            "\"details\":\"原文要点\",\"source_pages\":[1],\"confidence\":0.0}]}]}。"
            "source_pages 必须填写图片前标注的页码。目录、说明、题型、分值、范围都应保留为结构化要求。") }]
        for page_number, path in pages:
            content.extend([{"type": "text", "text": f"源文件页面：{page_number}"},
                            self._image_part(path.read_bytes())])
        return await self._vision_json(content, 8000)

    async def _merge_fragments(self, fragments: list[dict[str, Any]]) -> dict[str, Any]:
        if not fragments:
            raise UserSyllabusError(USER_SYLLABUS_INVALID_STRUCTURE, "未提取到考纲结构")
        merged = dict(fragments[0])
        sections: list[dict[str, Any]] = []
        seen: set[str] = set()
        for fragment in fragments:
            for section in fragment.get("sections") or []:
                key = str(section.get("title") or "").strip()
                if key and key in seen:
                    continue
                if key:
                    seen.add(key)
                sections.append(section)
        merged["sections"] = sections
        return merged

    def _mineru_markdown(self, pdf_path: Path, output_dir: Path) -> Path:
        if not self.mineru_token:
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "MinerU 服务端密钥未配置")
        if not self.mineru_pipeline_root:
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "MinerU 处理管线未配置")
        script = self.mineru_pipeline_root / "parse_question_pdf.py"
        config = self.mineru_pipeline_root / "pipeline_config.json"
        if not script.is_file() or not config.is_file():
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "MinerU 处理管线不完整")
        output_dir.mkdir(parents=True, exist_ok=True)
        env = os.environ.copy()
        env["MINERU_TOKEN"] = self.mineru_token
        command = [
            sys.executable, str(script),
            "--config", str(config), "--output-dir", str(output_dir),
            "--pdf", str(pdf_path),
        ]
        completed = None
        for attempt in range(3):
            completed = subprocess.run(
                command, cwd=self.mineru_pipeline_root, env=env,
                capture_output=True, text=True, encoding="utf-8", errors="replace",
                timeout=24 * 60 * 60, check=False,
            )
            if completed.returncode == 0:
                break
            detail = (completed.stderr or completed.stdout or "").lower()
            if "download failed after retries" not in detail or attempt >= 2:
                break
            time.sleep(4 * (attempt + 1))
        assert completed is not None
        if completed.returncode != 0:
            detail = (completed.stderr or completed.stdout or "MinerU 解析失败").strip()
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, detail[-1500:] or "MinerU 解析失败")
        markdown_files = sorted(output_dir.rglob("*_clean.md")) or sorted(output_dir.rglob("*.md"))
        if not markdown_files:
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "MinerU 未生成考纲 Markdown")
        combined = output_dir / "syllabus_full_clean.md"
        combined.write_text(
            "\n\n".join(path.read_text(encoding="utf-8-sig") for path in markdown_files),
            encoding="utf-8",
        )
        return combined

    async def _extract_markdown_batches(self, markdown: Path) -> list[dict[str, Any]]:
        text = markdown.read_text(encoding="utf-8", errors="replace").strip()
        if not text:
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "MinerU 未生成考纲内容")
        chunks = _split_text(text, 3000)
        semaphore = asyncio.Semaphore(6)

        async def extract(chunk: str) -> dict[str, Any]:
            async with semaphore:
                return await self._extract_markdown_chunk(chunk)

        results = await asyncio.gather(*(extract(chunk) for chunk in chunks), return_exceptions=True)
        fragments = [value for value in results if isinstance(value, dict)]
        if not fragments:
            failure = next((value for value in results if isinstance(value, Exception)), None)
            raise UserSyllabusError(
                USER_SYLLABUS_EXTRACTION_FAILED,
                f"考纲文本结构化失败：{type(failure).__name__}: {failure}" if failure else "考纲文本结构化失败"
            )
        return fragments

    async def _extract_markdown_chunk(self, text: str) -> dict[str, Any]:
        return await self._vision_json([{
            "type": "text",
            "text": (
                "你是考试考纲结构化解析器。以下内容是同一份考纲的 Markdown 文本，"
                "已标注源文件页码（如 第3页 或 <!-- page 3 -->）。"
                "只根据文本内容识别，不得补写文本中不存在的要求。返回严格 JSON："
                "{\"document_title\":\"\",\"subject\":\"\",\"exam_type\":\"\","
                "\"sections\":[{\"title\":\"章节/模块名\",\"requirements\":["
                "{\"title\":\"具体考核要求\",\"mastery_level\":\"掌握/熟悉/了解/未注明\","
                "\"details\":\"原文要点\",\"source_pages\":[1],\"confidence\":0.0}]}]}。"
                "source_pages 必须填写文本中标注的页码。目录、说明、题型、分值、范围都应保留为结构化要求。\n\n"
                + text
            ),
        }], 8000)

    async def _vision_json(self, user_content: list[dict[str, Any]], max_tokens: int) -> dict[str, Any]:
        payload: dict[str, Any] = {"model": self.chat_model, "messages": [
            {"role": "system", "content": "你只负责识别和结构化用户考纲，输出有效 JSON，不输出 Markdown。"},
            {"role": "user", "content": user_content}],
            "response_format": {"type": "json_object"}, "temperature": 0, "max_tokens": max_tokens,
            "reasoning_effort": "none"}
        timeout = httpx.Timeout(max(self.timeout_seconds, 600.0), connect=30.0)
        async with httpx.AsyncClient(timeout=timeout) as client:
            for attempt in range(3):
                try:
                    response = await client.post(f"{self.chat_base_url}/chat/completions",
                        headers={"Authorization": f"Bearer {self.chat_api_key}"}, json=payload)
                    if response.status_code == 400:
                        lowered = response.text.lower()
                        if "response_format" in lowered:
                            payload.pop("response_format", None)
                            response = await client.post(f"{self.chat_base_url}/chat/completions",
                                headers={"Authorization": f"Bearer {self.chat_api_key}"}, json=payload)
                        elif "reasoning_effort" in lowered:
                            payload.pop("reasoning_effort", None)
                            response = await client.post(f"{self.chat_base_url}/chat/completions",
                                headers={"Authorization": f"Bearer {self.chat_api_key}"}, json=payload)
                    if response.status_code in (408, 429, 500, 502, 503, 504):
                        if attempt < 2:
                            await asyncio.sleep(2 * (attempt + 1))
                            continue
                        raise UserSyllabusError(
                            USER_SYLLABUS_EXTRACTION_FAILED,
                            f"多模态模型调用失败（HTTP {response.status_code}）")
                    response.raise_for_status()
                    raw = response.json().get("choices", [{}])[0].get("message", {}).get("content", "")
                    if not str(raw or "").strip():
                        if attempt < 2:
                            await asyncio.sleep(2 * (attempt + 1))
                            continue
                        raise UserSyllabusError(
                            USER_SYLLABUS_EXTRACTION_FAILED, "多模态模型未返回内容（推理过长或参数被忽略）")
                    try:
                        return _json_object(raw)
                    except UserSyllabusError:
                        if attempt < 2:
                            await asyncio.sleep(2 * (attempt + 1))
                            continue
                        raise
                except UserSyllabusError:
                    raise
                except (httpx.HTTPStatusError, httpx.TransportError) as exc:
                    if attempt < 2:
                        await asyncio.sleep(2 * (attempt + 1))
                        continue
                    raise UserSyllabusError(
                        USER_SYLLABUS_EXTRACTION_FAILED,
                        f"多模态模型调用失败：{type(exc).__name__}: {exc}") from exc

    @staticmethod
    def _image_part(content: bytes) -> dict[str, Any]:
        data = base64.b64encode(content).decode("ascii")
        return {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{data}"}}

    def _render_input(self, source: Path, output_dir: Path) -> list[tuple[int, Path]]:
        if output_dir.exists():
            shutil.rmtree(output_dir)
        output_dir.mkdir(parents=True)
        suffix = source.suffix.lower()
        if suffix == ".pdf":
            return self._render_pdf(source, output_dir)
        if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
            return self._render_images(source, output_dir)
        if suffix == ".docx":
            document = Document(source)
            lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            pages = self._render_text_lines(lines, output_dir)
            for rel in document.part.rels.values():
                target = getattr(rel, "target_part", None)
                blob = getattr(target, "blob", None)
                if blob and str(getattr(target, "content_type", "")).startswith("image/"):
                    pages.extend(self._append_image_blob(blob, output_dir, len(pages) + 1))
            return pages
        if suffix == ".xlsx":
            workbook = load_workbook(source, read_only=True, data_only=True)
            lines = []
            for sheet in workbook.worksheets:
                lines.append(f"工作表：{sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value not in (None, "")]
                    if values:
                        lines.append(" | ".join(values))
            workbook.close()
            return self._render_text_lines(lines, output_dir)
        if suffix in {".doc", ".xls"}:
            converted = self._convert_office_to_pdf(source)
            try:
                return self._render_pdf(converted, output_dir)
            finally:
                shutil.rmtree(converted.parent, ignore_errors=True)
        if suffix in {".md", ".markdown", ".txt", ".csv"}:
            return self._render_text_lines(self._decode_text(source.read_bytes()).splitlines(), output_dir)
        raise UserSyllabusError(USER_SYLLABUS_UNSUPPORTED_FILE, "不支持该考纲文件格式")

    def _render_pdf(self, source: Path, output_dir: Path) -> list[tuple[int, Path]]:
        try:
            count = len(PdfReader(str(source)).pages)
        except Exception as exc:
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "PDF 无法读取") from exc
        pages = []
        for page in range(1, count + 1):
            prefix = output_dir / f"render_{page:04d}"
            completed = subprocess.run(["pdftoppm", "-f", str(page), "-l", str(page),
                "-png", "-r", "150", str(source), str(prefix)], capture_output=True,
                timeout=180, check=False)
            candidates = sorted(output_dir.glob(f"{prefix.name}-*.png"))
            if completed.returncode or not candidates:
                raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, f"PDF 第 {page} 页渲染失败")
            target = output_dir / f"page_{page:04d}.jpg"
            with Image.open(candidates[0]) as rendered:
                rendered.convert("RGB").save(target, "JPEG", quality=90)
            candidates[0].unlink(missing_ok=True)
            pages.append((page, target))
        return pages

    @staticmethod
    def _render_images(source: Path, output_dir: Path) -> list[tuple[int, Path]]:
        pages = []
        with Image.open(source) as image:
            for index, frame in enumerate(ImageSequence.Iterator(image), start=1):
                target = output_dir / f"page_{index:04d}.jpg"
                frame.convert("RGB").save(target, "JPEG", quality=90)
                pages.append((index, target))
        return pages

    @staticmethod
    def _append_image_blob(blob: bytes, output_dir: Path, page: int) -> list[tuple[int, Path]]:
        try:
            with Image.open(BytesIO(blob)) as image:
                target = output_dir / f"page_{page:04d}.jpg"
                image.convert("RGB").save(target, "JPEG", quality=90)
                return [(page, target)]
        except Exception:
            return []

    def _render_text_lines(self, lines: list[str], output_dir: Path) -> list[tuple[int, Path]]:
        normalized = []
        for line in lines:
            value = str(line or "").rstrip()
            normalized.extend(textwrap.wrap(value, width=52, break_long_words=True,
                                             break_on_hyphens=False) or [""])
        if not any(line.strip() for line in normalized):
            return []
        font, label_font = self._font(28), self._font(22)
        pages = []
        for offset in range(0, len(normalized), 34):
            page = len(pages) + 1
            image = Image.new("RGB", (1240, 1754), "white")
            draw = ImageDraw.Draw(image)
            draw.text((70, 40), f"源文件虚拟页面 {page}", fill="#64748b", font=label_font)
            for row, line in enumerate(normalized[offset:offset + 34]):
                draw.text((70, 100 + row * 46), line, fill="#111827", font=font)
            target = output_dir / f"page_{page:04d}.jpg"
            image.save(target, "JPEG", quality=90)
            pages.append((page, target))
        return pages

    @staticmethod
    def _font(size: int):
        for candidate in (Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simhei.ttf"),
                          Path("C:/Windows/Fonts/simsun.ttc")):
            if candidate.is_file():
                return ImageFont.truetype(str(candidate), size=size)
        return ImageFont.load_default()

    @staticmethod
    def _decode_text(content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                pass
        return content.decode("utf-8", errors="replace")

    @staticmethod
    def _convert_office_to_pdf(source: Path) -> Path:
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED,
                                    "旧版 Word/Excel 需要安装 LibreOffice 后才能转换")
        output_dir = Path(tempfile.mkdtemp(prefix="user-syllabus-office-"))
        completed = subprocess.run([soffice, "--headless", "--convert-to", "pdf",
            "--outdir", str(output_dir), str(source)], capture_output=True, timeout=180, check=False)
        converted = output_dir / f"{source.stem}.pdf"
        if completed.returncode or not converted.is_file():
            shutil.rmtree(output_dir, ignore_errors=True)
            raise UserSyllabusError(USER_SYLLABUS_EXTRACTION_FAILED, "旧版 Word/Excel 转换失败")
        return converted

    def _normalize_structure(self, raw: dict[str, Any], syllabus_id: str, fallback_title: str,
                             fallback_subject: str, fallback_exam_type: str
                             ) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        sections_raw = raw.get("sections")
        if not isinstance(sections_raw, list):
            raise UserSyllabusError(USER_SYLLABUS_INVALID_STRUCTURE, "多模态模型未返回考纲章节")
        sections, requirements = [], []
        for section_index, section_raw in enumerate(sections_raw, start=1):
            if not isinstance(section_raw, dict):
                continue
            section_title = str(section_raw.get("title") or f"第{section_index}部分").strip()
            section_id = f"SEC_{syllabus_id[4:16]}_{section_index:03d}"
            section_requirements = []
            raw_requirements = section_raw.get("requirements") or section_raw.get("children") or []
            if not isinstance(raw_requirements, list):
                continue
            for item in raw_requirements:
                if not isinstance(item, dict):
                    continue
                title = str(item.get("title") or item.get("requirement") or item.get("content") or "").strip()
                if not title:
                    continue
                pages = []
                for value in item.get("source_pages") or []:
                    try:
                        page = int(value)
                    except (TypeError, ValueError):
                        continue
                    if page > 0 and page not in pages:
                        pages.append(page)
                try:
                    confidence = min(1.0, max(0.0, float(item.get("confidence", 0))))
                except (TypeError, ValueError):
                    confidence = 0.0
                requirement = {
                    "requirement_id": f"REQ_{uuid4().hex}", "section_id": section_id,
                    "section_title": section_title, "title": title,
                    "mastery_level": str(item.get("mastery_level") or "未注明").strip(),
                    "details": str(item.get("details") or "").strip(),
                    "source_pages": pages, "confidence": confidence,
                }
                requirements.append(requirement)
                section_requirements.append(requirement)
            if section_requirements:
                sections.append({"section_id": section_id, "title": section_title,
                                 "requirements": section_requirements})
        if not requirements:
            raise UserSyllabusError(USER_SYLLABUS_INVALID_STRUCTURE, "未识别到任何考纲要求")
        structured = {
            "syllabus_id": syllabus_id, "scope": "personal", "origin": "user_upload",
            "title": str(raw.get("document_title") or raw.get("title") or fallback_title).strip(),
            "subject": str(raw.get("subject") or fallback_subject).strip(),
            "exam_type": str(raw.get("exam_type") or fallback_exam_type).strip(),
            "sections": sections,
        }
        return structured, requirements

    def _match_requirement(self, resolver: Any, requirement: dict[str, Any],
                           structured: dict[str, Any]) -> dict[str, Any]:
        query = "；".join(filter(None, [structured.get("subject"), requirement.get("section_title"),
                                         requirement.get("title"), requirement.get("details")]))
        candidates = []
        try:
            raw_candidates = resolver(query, limit=3)
        except Exception:
            raw_candidates = []
        for candidate in raw_candidates or []:
            if hasattr(candidate, "model_dump"):
                candidate = candidate.model_dump(mode="json")
            if isinstance(candidate, dict):
                candidates.append({"kp_id": str(candidate.get("kp_id") or ""),
                    "kp_name": str(candidate.get("name") or candidate.get("kp_name") or ""),
                    "confidence": float(candidate.get("score") or candidate.get("confidence") or 0)})
        best = candidates[0] if candidates and candidates[0]["confidence"] >= 0.55 else None
        return {"requirement_id": requirement["requirement_id"],
            "match_status": "matched" if best else "unmatched",
            "kp_id": best["kp_id"] if best else None,
            "kp_name": best["kp_name"] if best else None,
            "confidence": best["confidence"] if best else 0.0,
            "match_source": "public_kp_retrieval" if best else None,
            "candidates": candidates}

    async def _map_requirements(self, requirements: list[dict[str, Any]],
                                structured: dict[str, Any]) -> list[dict[str, Any]]:
        if self.vector_matcher is not None:
            try:
                return await self.vector_matcher.match(requirements, structured)
            except Exception as exc:
                self.vector_matcher = None
                if getattr(self, "_vector_match_error", None) is None:
                    self._vector_match_error = f"{type(exc).__name__}: {exc}"
        return self._map_requirements_lexical(requirements, structured)

    def _map_requirements_lexical(self, requirements: list[dict[str, Any]],
                                  structured: dict[str, Any]) -> list[dict[str, Any]]:
        resolver = getattr(self.knowledge_resolver, "resolve_topic", None)
        if not resolver:
            return [self._match_requirement(None, requirement, structured)
                    for requirement in requirements]
        owner = getattr(resolver, "__self__", None)
        if owner is not None:
            for name in ("ensure_hierarchy", "_kp_search_entries"):
                warmup = getattr(owner, name, None)
                if callable(warmup):
                    try:
                        warmup()
                    except Exception:
                        pass
        with ThreadPoolExecutor(max_workers=8) as pool:
            return list(pool.map(
                lambda requirement: self._match_requirement(resolver, requirement, structured),
                requirements,
            ))

    def _syllabus_dir(self, owner: str, syllabus_id: str) -> Path:
        if not re.fullmatch(r"USY_[A-Za-z0-9_-]{8,80}", str(syllabus_id or "")):
            raise UserSyllabusError(USER_SYLLABUS_NOT_FOUND, "考纲不存在")
        return self.root / owner / syllabus_id

    def _require_dir(self, owner: str, syllabus_id: str) -> Path:
        path = self._syllabus_dir(owner, syllabus_id)
        if not path.is_dir() or not (path / "manifest.json").is_file():
            raise UserSyllabusError(USER_SYLLABUS_NOT_FOUND, "考纲不存在")
        return path

    @staticmethod
    def _read_json(path: Path, fallback: Any) -> Any:
        if not path.is_file():
            return fallback
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return fallback

    @staticmethod
    def _read_jsonl(path: Path) -> list[dict[str, Any]]:
        if not path.is_file():
            return []
        rows = []
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                try:
                    value = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if isinstance(value, dict):
                    rows.append(value)
        return rows

    @staticmethod
    def _write_json(path: Path, value: Any) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(path.suffix + ".tmp")
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)

    @staticmethod
    def _write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(path.suffix + ".tmp")
        temporary.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
                             encoding="utf-8")
        temporary.replace(path)
